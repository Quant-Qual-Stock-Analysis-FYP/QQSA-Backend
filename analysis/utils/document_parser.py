"""Utility functions for parsing document files (docx, pdf) into text."""

import logging
from io import BytesIO
from typing import Optional

logger = logging.getLogger(__name__)


def extract_text_from_file(file, filename: str) -> Optional[str]:
    """
    Extract text content from uploaded file (docx or pdf).
    
    Args:
        file: Django uploaded file object
        filename: Original filename
        
    Returns:
        Extracted text content as string, or None if extraction fails
    """
    try:
        # Get file extension
        file_ext = filename.lower().split('.')[-1] if '.' in filename else ''
        
        if file_ext == 'docx':
            return _extract_text_from_docx(file)
        elif file_ext == 'pdf':
            return _extract_text_from_pdf(file)
        else:
            logger.error(f"Unsupported file type: {file_ext}")
            return None
            
    except Exception as e:
        logger.error(f"Error extracting text from file {filename}: {e}", exc_info=True)
        return None


def _extract_text_from_docx(file) -> str:
    """Extract text from a .docx file."""
    try:
        from docx import Document
        
        # Read file content
        file.seek(0)  # Reset file pointer
        doc = Document(BytesIO(file.read()))
        
        # Extract text from all paragraphs
        text_parts = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text.strip())
        
        # Also extract text from tables if any
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_text.append(cell.text.strip())
                if row_text:
                    text_parts.append(" | ".join(row_text))
        
        return "\n\n".join(text_parts)
        
    except ImportError:
        logger.error("python-docx library not installed")
        raise
    except Exception as e:
        logger.error(f"Error parsing docx file: {e}", exc_info=True)
        raise


def _extract_text_from_pdf(file) -> str:
    """Extract text from a .pdf file."""
    try:
        from pypdf import PdfReader
        
        # Read file content
        file.seek(0)  # Reset file pointer
        pdf_reader = PdfReader(BytesIO(file.read()))
        
        # Extract text from all pages
        text_parts = []
        for page in pdf_reader.pages:
            page_text = page.extract_text()
            if page_text.strip():
                text_parts.append(page_text.strip())
        
        return "\n\n".join(text_parts)
        
    except ImportError:
        logger.error("pypdf library not installed")
        raise
    except Exception as e:
        logger.error(f"Error parsing pdf file: {e}", exc_info=True)
        raise
